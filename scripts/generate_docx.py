"""Genera qa-challenge-reto[nombre].docx con respuestas de los retos 1, 3, 4 y 5."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = r"C:\Users\user\Projects\EvaluacionTecnicaQA\qa-challenge-reto-candidato.docx"


def heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    return p


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()


def main():
    doc = Document()
    title = doc.add_heading("Evaluación Técnica QA — Respuestas", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(
        doc,
        "Nota: Renombrar este archivo a qa-challenge-reto[TuNombre].docx antes de enviarlo.",
        bold=True,
    )
    doc.add_paragraph()

    # RETO 1
    heading(doc, "RETO 1 — Control de calidad (Historia de usuario: Registro y Login)", 1)

    heading(doc, "1. Actividades para asegurar un óptimo proceso de control de calidad", 2)
    bullets(
        doc,
        [
            "Análisis de requerimientos y criterios de aceptación: revisar la HU con el PO, aclarar ambigüedades (¿5 caracteres inclusive?, mensajes de error esperados, alcance front/back) y obtener criterios medibles.",
            "Revisión estática del código y contratos API (endpoints, códigos HTTP, payloads) para alinear pruebas de servicio y UI antes de ejecutar.",
            "Diseño de casos de prueba (funcionales, negativos, límites) con trazabilidad requisito → caso → evidencia.",
            "Preparación de datos y entorno: levantar front (4000) y back (3000), definir usuarios/correos de prueba y checklist de smoke.",
            "Ejecución manual exploratoria + regresión de casos diseñados; registro de defectos con severidad, pasos, esperado vs actual y capturas.",
            "Automatización de casos críticos (regresión) en pipeline CI cuando el código esté estable.",
            "Reporte de calidad al equipo: resumen de cobertura, riesgos, defectos abiertos y recomendación de release (Go/No-Go).",
        ],
    )
    para(
        doc,
        "¿Por qué? Porque reduce riesgo de escape a producción, da visibilidad temprana de desalineación con criterios de aceptación y documenta decisiones para auditoría y mejora continua.",
    )

    heading(doc, "2. Tipos de prueba aplicables", 2)
    bullets(
        doc,
        [
            "Pruebas funcionales: validar registro y login según criterios de aceptación.",
            "Pruebas de API/servicios: POST /register y /login (validaciones, códigos HTTP, mensajes).",
            "Pruebas de UI/E2E: flujos desde el navegador hacia el backend.",
            "Pruebas negativas y de validación: campos vacíos, email sin @, contraseña corta.",
            "Pruebas de límites (boundary): contraseña con exactamente 5 caracteres.",
            "Pruebas de integración front-back: CORS, manejo de errores en axios.",
            "Smoke/Sanity tras cada despliegue en QA.",
            "Regresión automatizada sobre casos críticos (Reto 2).",
        ],
    )

    heading(doc, "3. Técnicas de diseño de casos de prueba (≥4)", 2)
    bullets(
        doc,
        [
            "Clases de equivalencia: particionar email (válido con @, sin @, vacío) y contraseña (≥5, <5, vacía) para reducir casos redundantes.",
            "Análisis de valores límite: probar 4, 5 y 6 caracteres en contraseña; email con un solo carácter antes/después del @.",
            "Tabla de decisión: combinaciones email/contraseña presentes o ausentes → REQUIRED vs INVALID vs éxito.",
            "Pruebas basadas en estados: usuario no registrado vs registrado (login); hoy el sistema no persiste usuarios — se documenta como gap.",
            "Pairwise (opcional): combinar tipos de email inválido con longitudes de contraseña para cobertura eficiente.",
        ],
    )

    heading(doc, "4. Matriz de casos de prueba y orden de ejecución", 2)
    add_table(
        doc,
        ["ID", "Módulo", "Caso", "Prioridad", "Orden", "Motivo del orden"],
        [
            ["TC-01", "API Registro", "Registro exitoso email@dominio + pass ≥5", "Alta", "1", "Smoke: confirma que el flujo feliz y el entorno responden."],
            ["TC-02", "API Registro", "Campos vacíos → REQUIRED", "Alta", "2", "Valida regla obligatoria antes de formatos."],
            ["TC-03", "API Registro", "Email sin @ → INVALID", "Alta", "3", "Regla de formato crítica."],
            ["TC-04", "API Registro", "Contraseña 4 caracteres → INVALID", "Alta", "4", "Límite inferior inválido."],
            ["TC-05", "API Registro", "Contraseña exactamente 5 caracteres", "Alta", "5", "Boundary; verifica criterio 'mínimo 5'."],
            ["TC-06", "UI Registro", "Mismos TC-02 a TC-05 en pantalla /register", "Alta", "6", "Cobertura front + mensaje en #msg."],
            ["TC-07", "API Login", "Login con credenciales inválidas → error", "Alta", "7", "Seguridad y criterios de login."],
            ["TC-08", "API Login", "Email sin @ en login → INVALID", "Alta", "8", "Paridad con registro."],
            ["TC-09", "API Login", "Contraseña <5 en login → INVALID", "Alta", "9", "Paridad con registro."],
            ["TC-10", "UI Login", "Casos negativos y éxito en /login", "Media", "10", "E2E completo tras validar API."],
        ],
    )
    para(
        doc,
        "Orden: primero API (aislar capa de negocio), luego UI (depende de API y CORS), de lo general (obligatoriedad) a lo específico (límites), y smoke antes de regresión profunda.",
    )

    heading(doc, "5. Criterios para estimar story points (nuevo en el equipo)", 2)
    bullets(
        doc,
        [
            "Complejidad funcional: dos flujos (registro + login) con validaciones duplicadas front/back.",
            "Incertidumbre: contratos API, persistencia de usuarios y mensajes UX no detallados → spike o puntos extra.",
            "Dependencias: entornos, datos, posible BD/auth real no implementada.",
            "Riesgo/regulación: manejo de credenciales (aunque sea demo, implica seguridad).",
            "Esfuerzo de prueba: diseño de matriz, ejecución manual, automatización inicial, evidencias.",
            "Comparación con historias referencia del equipo (planning poker) tras explicar supuestos.",
        ],
    )
    para(doc, "Estimación orientativa: 5–8 SP si incluye automatización básica y certificación front+back; 3–5 SP si solo manual con entorno ya estable.")

    heading(doc, "6. Bugs y mejoras (≥5 bugs y ≥3 mejoras por requerimiento)", 2)
    heading(doc, "Registro", 3)
    add_table(
        doc,
        ["ID", "Tipo", "Severidad", "Descripción", "Esperado", "Actual"],
        [
            ["BUG-R01", "Bug", "Alta", "Longitud mínima de contraseña", "Aceptar 5 caracteres", "422 INVALID con 5 chars (usa >5 en server.js)"],
            ["BUG-R02", "Bug", "Alta", "Validación solo en backend", "Validar en front y back", "Front envía sin validar localmente"],
            ["BUG-R03", "Bug", "Media", "Mensajes de error", "Mensaje claro al usuario", "REQUIRED / INVALID genéricos"],
            ["BUG-R04", "Bug", "Media", "Email débil", "Email válido con @", "Acepta '@' o '@@' como válido"],
            ["BUG-R05", "Bug", "Alta", "Persistencia", "Usuario guardado para login", "Siempre SAVED sin almacenamiento"],
            ["BUG-R06", "Bug", "Baja", "Atributos HTML", "Campos required / type=email", "type=text sin required"],
            ["IMP-R01", "Mejora", "-", "i18n mensajes en español", "-", "-"],
            ["IMP-R02", "Mejora", "-", "Hash de contraseña (bcrypt) y BD", "-", "-"],
            ["IMP-R03", "Mejora", "-", "Feedback UX: loading y deshabilitar botón", "-", "-"],
        ],
    )
    heading(doc, "Login", 3)
    add_table(
        doc,
        ["ID", "Tipo", "Severidad", "Descripción", "Esperado", "Actual"],
        [
            ["BUG-L01", "Bug", "Crítica", "Autenticación", "Rechazar credenciales incorrectas", "LOGIN VALID con cualquier email/pass no vacío"],
            ["BUG-L02", "Bug", "Alta", "Validación email en login", "Rechazar sin @", "No valida formato en /login"],
            ["BUG-L03", "Bug", "Alta", "Validación contraseña en login", "Mínimo 5 caracteres", "Acepta pass de 3 caracteres"],
            ["BUG-L04", "Bug", "Alta", "Validación front login", "Errores antes de API", "Sin validación cliente"],
            ["BUG-L05", "Bug", "Media", "Código comentado", "Lógica clara de auth", "Condición de login comentada en server.js"],
            ["BUG-L06", "Bug", "Baja", "Código HTTP", "200 en login OK", "201 LOGIN VALID"],
            ["IMP-L01", "Mejora", "-", "Sesión/JWT tras login exitoso", "-", "-"],
            ["IMP-L02", "Mejora", "-", "Mensajes específicos por campo", "-", "-"],
            ["IMP-L03", "Mejora", "-", "Rate limiting / bloqueo por intentos", "-", "-"],
        ],
    )

    heading(doc, "7. Desarrollador indica que 2 bugs no aplican", 2)
    bullets(
        doc,
        [
            "Revisar el criterio de aceptación y la HU con evidencia (capturas, request/response).",
            "Sesión breve con dev + PO para alinear alcance; si el PO confirma que no aplica, cerrar como 'No defecto / Cambio de alcance' documentado.",
            "Si sigue en alcance, escalar con impacto en negocio/seguridad y proponer fix o waiver formal con fecha.",
        ],
    )

    heading(doc, "8. Bloqueo por dependencias (accesos, data, persona no disponible)", 2)
    bullets(
        doc,
        [
            "Registrar el bloqueante en el tablero con impacto en fechas.",
            "Usar mocks/stubs locales (datos de prueba, Postman collections) para avanzar en API/UI.",
            "Pedir acceso alterno o ambiente compartido documentado; escalar al TL/PO.",
            "Ejecutar lo posible (casos no bloqueados) y comunicar riesgo residual.",
        ],
    )

    heading(doc, "9. PO solicita cambios durante la ejecución", 2)
    bullets(
        doc,
        [
            "Congelar alcance de la iteración actual o negociar swap de scope.",
            "Impact analysis: casos afectados, regresión, estimación adicional.",
            "Actualizar matriz de prueba y criterios de aceptación por escrito antes de continuar.",
            "Re-priorizar automatización según nuevos riesgos.",
        ],
    )

    heading(doc, "10. Fecha límite con presión de entrega", 2)
    bullets(
        doc,
        [
            "Comunicar estado transparente: qué está probado, qué no, defectos abiertos críticos.",
            "Proponer release con riesgo documentado o reducción de alcance acordada con PO.",
            "Priorizar pruebas de riesgo alto (seguridad, flujos bloqueantes); diferir mejoras menores.",
            "No omitir evidencia de defectos críticos por presión; buscar extensión mínima o hardening en hotfix.",
        ],
    )

    # RETO 3
    heading(doc, "RETO 3 — Niveles de prueba por ambiente", 1)
    add_table(
        doc,
        ["Ambiente", "Niveles recomendados", "Comentario"],
        [
            ["Desarrollo (DEV)", "Unitarias, componentes, contratos API locales", "Feedback rápido al dev; mocks de servicios externos."],
            ["QA / Test", "Integración, API, E2E automatizada, regresión, exploratoria", "Ambiente estable; datos sintéticos; CI nightly."],
            ["UAT", "E2E manuales guiados, aceptación usuario, smoke negocio", "Validación con PO/usuarios clave; datos representativos."],
            ["Pre-producción / Staging", "Smoke, regresión completa, performance ligera, seguridad básica", "Paridad con prod; prueba de despliegue."],
            ["Producción", "Smoke post-deploy, monitoreo, pruebas sintéticas (canary)", "Sin pruebas destructivas; observabilidad."],
        ],
    )
    para(
        doc,
        "La pirámide sugiere muchas unitarias/integración en DEV, menos E2E costosos concentrados en QA/UAT/Staging. Producción no sustituye QA pero valida salud del servicio.",
    )

    # RETO 4
    heading(doc, "RETO 4 — Priorización de defectos (1 = más crítico)", 1)
    add_table(
        doc,
        ["Prioridad", "Defecto", "Justificación breve"],
        [
            ["1", "Credenciales sin cifrar en BD", "Riesgo de seguridad y cumplimiento; exposición masiva de cuentas."],
            ["2", "Precios no varían entre monedas", "Impacto directo en dinero del usuario y confianza en compraventa global."],
            ["3", "Transacciones no en informe back-end", "Integridad contable/operativa; decisiones de negocio con datos incorrectos."],
            ["4", "Ancho de banda 100 MB vs objetivo 50 MB", "Degrada experiencia en mercados con conectividad limitada; contradice propuesta de producto."],
            ["5", "Foros no registran comentarios", "Soporte degradado; no bloquea compra pero afecta retención."],
            ["6", "Errores ortográficos en home", "Impacto cosmético en percepción de marca; bajo riesgo funcional."],
        ],
    )
    para(
        doc,
        "Supuestos: plataforma de comercio con pagos reales; cumplimiento de seguridad es no negociable; audiencia internacional.",
    )

    # RETO 5
    heading(doc, "RETO 5 — Evaluación de procedimientos Agile Testing", 1)
    procedures = [
        (
            "Performance en el mismo entorno que usuarios finales",
            "No ideal: mezcla tráfico real con carga de prueba; riesgo y resultados no reproducibles.",
            "Ambiente dedicado de performance (staging) con perfil de carga representativo; pruebas de humo de perf en CI con umbrales; monitoreo APM en prod sin load test invasivo.",
        ),
        (
            "Smoke manual cubriendo todas las opciones en cada deploy",
            "Smoke debe ser rápido y crítico; cubrir 'todo' no es smoke y no escala.",
            "Smoke automatizado (15–30 min) de rutas críticas: login, búsqueda, publicar, checkout; exploratoria manual en releases mayores.",
        ),
        (
            "Regresión solo en módulos 'probablemente afectados'",
            "Riesgoso por dependencias ocultas y deuda.",
            "Regresión por capas: suite automatizada core en cada PR; regresión completa en release candidate; análisis de impacto como complemento, no único filtro.",
        ),
        (
            "Casos en Excel eliminados cada versión",
            "Pérdida de conocimiento y retrabajo.",
            "Repositorio de casos en herramienta de gestión (Xray, TestRail, o repo Git); versionado; reutilización en automatización.",
        ),
        (
            "QA involucrado desde etapa inicial",
            "Adecuado: shift-left, revisión de criterios, TDD/BDD donde aplique.",
            "Mantener y reforzar con tres amigos, DoD con criterios testeables y gates de CI.",
        ),
    ]
    for i, (proc, eval_, alt) in enumerate(procedures, 1):
        heading(doc, f"Procedimiento {i}", 3)
        para(doc, f"Práctica: {proc}")
        para(doc, f"Evaluación: {eval_}")
        para(doc, f"Alternativa / mejora: {alt}")

    doc.save(OUTPUT)
    print(f"Generado: {OUTPUT}")


if __name__ == "__main__":
    main()
